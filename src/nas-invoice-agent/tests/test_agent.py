import io
import os
import time
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from invoice_agent import config, main, ocr, scanner, sheet
from invoice_agent.extract import ExtractionError, Invoice
from invoice_agent.state import State

INVOICE_TEXT = "TAX INVOICE\nAcme Supplies Pty Ltd\nInvoice No: INV-1042\nDate: 05/09/2026\nTotal: $110.00"


def make_text_pdf(path: Path, text: str = INVOICE_TEXT) -> Path:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=12)
    doc.save(path)
    return path


def make_scanned_pdf(path: Path, text: str = INVOICE_TEXT) -> Path:
    """Image-only PDF, like a scanner produces: no text layer, so OCR is required."""
    img = Image.new("RGB", (1700, 2200), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 48)
    except OSError:
        font = ImageFont.load_default(size=48)
    draw.multiline_text((120, 120), text, fill="black", font=font, spacing=24)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(path)
    return path


def age(path: Path, seconds: int = 3600) -> None:
    t = time.time() - seconds
    os.utime(path, (t, t))


class FakeExtractor:
    def __init__(self, fail_on: str | None = None):
        self.calls = []
        self.fail_on = fail_on

    def extract(self, text, filename):
        self.calls.append((filename, text))
        if self.fail_on and self.fail_on in filename:
            raise ExtractionError("boom")
        return Invoice(
            document_type="invoice", supplier_name="Acme Supplies Pty Ltd", supplier_abn=None,
            invoice_number="INV-1042", invoice_date="2026-09-05", due_date=None, customer_name=None,
            currency="AUD", subtotal=100.0, tax_amount=10.0, total=110.0, description="Stationery",
            confidence="high", notes=None,
        )


@pytest.fixture
def settings(tmp_path, monkeypatch):
    monkeypatch.setenv("SCAN_PATHS", str(tmp_path / "scan"))
    monkeypatch.setenv("SPREADSHEET_PATH", str(tmp_path / "out" / "master.xlsx"))
    monkeypatch.setenv("STATE_DB_PATH", str(tmp_path / "state.db"))
    (tmp_path / "scan").mkdir()
    return config.load()


def test_text_layer_is_used_without_ocr(tmp_path):
    pdf = ocr.extract_text(make_text_pdf(tmp_path / "a.pdf"), 10, 300, "eng", 20)
    assert "INV-1042" in pdf.text
    assert pdf.ocr_pages == 0


def test_scanned_pdf_is_ocrd(tmp_path):
    pdf = ocr.extract_text(make_scanned_pdf(tmp_path / "scan.pdf"), 10, 300, "eng", 20)
    assert pdf.ocr_pages == 1
    assert "INV-1042" in pdf.text
    assert "Acme" in pdf.text


def test_scanner_skips_excluded_dirs_and_fresh_files(tmp_path):
    root = tmp_path / "scan"
    (root / "2026" / "Sept").mkdir(parents=True)
    (root / "@eaDir").mkdir()
    old = make_text_pdf(root / "2026" / "Sept" / "old.PDF")
    age(old)
    age(make_text_pdf(root / "@eaDir" / "thumb.pdf"))
    make_text_pdf(root / "fresh.pdf")  # just written, may still be copying
    (root / "notes.txt").write_text("x")
    found = list(scanner.find_pdfs([root], ["@eaDir"], min_age_seconds=60))
    assert found == [old]


def test_end_to_end_writes_once_and_skips_duplicates(settings, tmp_path):
    scan = tmp_path / "scan"
    age(make_text_pdf(scan / "inv1.pdf"))
    age(make_scanned_pdf(scan / "inv2.pdf", INVOICE_TEXT + "\nRef 2"))
    dup = scan / "copy-of-inv1.pdf"
    dup.write_bytes((scan / "inv1.pdf").read_bytes())
    age(dup)

    state, ex = State(settings.state_db_path), FakeExtractor()
    main.run_once(settings, state, ex)
    main.run_once(settings, state, ex)  # second scan: nothing new

    assert len(ex.calls) == 2  # duplicate content sent once, nothing re-sent on rescan
    ws = load_workbook(settings.spreadsheet_path)["Invoices"]
    assert ws.max_row == 3  # header + 2
    assert ws.cell(row=1, column=5).value == "Supplier"
    assert ws.cell(row=2, column=5).value == "Acme Supplies Pty Ltd"


def test_locked_spreadsheet_queues_rows(settings, tmp_path):
    age(make_text_pdf(tmp_path / "scan" / "inv1.pdf"))
    out = settings.spreadsheet_path
    out.parent.mkdir(parents=True)
    lock = out.with_name("~$" + out.name)
    lock.write_text("")

    state = State(settings.state_db_path)
    main.run_once(settings, state, FakeExtractor())
    assert not out.exists()
    assert len(state.pending_rows()) == 1

    lock.unlink()
    main.run_once(settings, state, FakeExtractor())
    assert load_workbook(out)["Invoices"].max_row == 2
    assert state.pending_rows() == []


def test_errors_are_retried_then_given_up(settings, tmp_path):
    age(make_text_pdf(tmp_path / "scan" / "bad.pdf"))
    state, ex = State(settings.state_db_path), FakeExtractor(fail_on="bad")
    for _ in range(5):
        main.run_once(settings, state, ex)
    assert len(ex.calls) == settings.max_attempts
    assert state.errors()[0][1] == "boom"


def test_review_flag_on_totals_mismatch(tmp_path):
    inv = FakeExtractor().extract("", "x.pdf").model_copy(update={"total": 999.0})
    pdf = ocr.PdfText(text="", page_count=1, pages_read=1, ocr_pages=0)
    row = main.build_row(tmp_path / "x.pdf", "a" * 64, pdf, inv)
    assert row["needs_review"] == "YES"
    assert "does not equal total" in row["notes"]
